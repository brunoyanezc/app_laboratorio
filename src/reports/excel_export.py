from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.styles import PatternFill
from openpyxl.styles import Alignment

from reports.pvgc2_report import (
    build_pvgc2_report
)

from reports.summary_report import (
    build_summary_report
)

from config.profile_loader import (
    load_profile
)

from reports.pvgc2_100g_report import (
    build_pvgc2_100g_report
)

def autofit_columns(ws):

    for column in ws.columns:

        max_length = 0

        column_letter = column[0].column_letter

        for cell in column:

            try:

                if cell.value is not None:

                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )

            except Exception:
                pass

        ws.column_dimensions[
            column_letter
        ].width = max_length + 2


def style_header_row(
    ws,
    row_number
):

    fill = PatternFill(
        fill_type="solid",
        fgColor="D9D9D9"
    )

    for cell in ws[row_number]:

        cell.font = Font(
            bold=True
        )

        cell.fill = fill

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


def export_excel(
    results_df,
    filename="Reporte_GCFID.xlsx"
):

    profile = load_profile(
        "PVGC2"
    )

    compounds = profile[
        "report_order"
    ]

    pvgc2 = build_pvgc2_report(
        results_df
    )

    summary = build_summary_report(
        results_df
    )

    report_100g = build_pvgc2_100g_report(
    results_df
    )

    wb = Workbook()

    # ----------------------------------
    # HOJA PRINCIPAL
    # ----------------------------------

    ws = wb.active

    ws.title = (
        "AFOCR_AC_GRASOS_A (CCLAS)"
    )

    row1 = [
        "Sample",
        "Wt Sample (g)",
        "Wt C23 (g)"
    ]

    for compound in compounds:

        row1.append(compound)
        row1.append(compound)

    ws.append(row1)

    row2 = [
        "",
        "",
        ""
    ]

    for _ in compounds:

        row2.append("%Area")
        row2.append("Area")

    ws.append(row2)

    for _, row in pvgc2.iterrows():

        output = [
            row["sample_id"],
            row["sample_weight_g"],
            row["mass_c23_added_g"]
        ]

        for compound in compounds:

            output.append(
                row[f"{compound}_%Area"]
            )

            output.append(
                row[f"{compound}_Area"]
            )

        ws.append(output)

    style_header_row(ws, 1)
    style_header_row(ws, 2)

    ws.freeze_panes = "D3"

    ws.auto_filter.ref = ws.dimensions

    autofit_columns(ws)

        # ----------------------------------
    # HOJA g100g
    # ----------------------------------

    ws100 = wb.create_sheet(
        "AFOCR_AC_GRASOS_100G (CCLAS)"
    )

    ws100.append(
        list(report_100g.columns)
    )

    for _, row in report_100g.iterrows():

        ws100.append(
            list(row)
        )

    style_header_row(
        ws100,
        1
    )

    autofit_columns(
        ws100
    )

    # ----------------------------------
    # RESUMEN
    # ----------------------------------

    ws2 = wb.create_sheet(
        "Resumen"
    )

    ws2.append(
        list(summary.columns)
    )

    for _, row in summary.iterrows():

        ws2.append(
            list(row)
        )

    style_header_row(ws2, 1)

    autofit_columns(ws2)

    # ----------------------------------
    # TRAZABILIDAD
    # ----------------------------------

    ws3 = wb.create_sheet(
        "Trazabilidad"
    )

    ws3.append(
        list(results_df.columns)
    )

    for _, row in results_df.iterrows():

        ws3.append(
            list(row)
        )

    style_header_row(ws3, 1)

    autofit_columns(ws3)

    wb.save(filename)

    return filename